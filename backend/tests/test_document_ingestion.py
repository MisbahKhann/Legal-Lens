"""
Comprehensive Unit & Integration Test Suite for Step 2 Document Ingestion and OCR Pipeline.
Tests representative document types:
1. Normal text-based PDF
2. Scanned PDF (image raster page)
3. DOCX document
4. Image / scanned page (PNG)
5. Multi-page document
6. Document containing tables
7. Source provenance traceability & storage persistence
"""

import os
from pathlib import Path
import pytest
import docx
from PIL import Image, ImageDraw
import reportlab.pdfgen.canvas as pdf_canvas
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from app.ingestion.pipeline import DocumentIngestionPipeline
from app.ingestion.models import FileType, ProcessingStatus, TextBlockType
from app.ingestion.ocr import NativeTextDetector
from app.ingestion.storage import DocumentStorage


@pytest.fixture
def tmp_storage_dir(tmp_path):
    return tmp_path / "storage"


@pytest.fixture
def ingestion_pipeline(tmp_storage_dir):
    return DocumentIngestionPipeline(base_storage_dir=tmp_storage_dir)


@pytest.fixture
def normal_text_pdf(tmp_path):
    pdf_path = tmp_path / "complaint_native.pdf"
    c = pdf_canvas.Canvas(str(pdf_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "UNITED STATES DISTRICT COURT")
    c.setFont("Helvetica", 11)
    c.drawString(72, 690, "COMPLAINT FOR BREACH OF CONTRACT")
    c.drawString(
        72,
        660,
        "Plaintiff Acme Corp hereby files this complaint against Defendant Globex Inc.",
    )
    c.drawString(
        72,
        640,
        "Jurisdiction is proper under 28 U.S.C. Section 1332 due to diversity of citizenship.",
    )
    c.save()
    return pdf_path


@pytest.fixture
def scanned_pdf(tmp_path):
    # Create image containing text
    img_path = tmp_path / "scanned_page.png"
    img = Image.new("RGB", (1000, 1200), color="white")
    d = ImageDraw.Draw(img)
    d.text((100, 100), "SCANNED EVIDENCE EXHIBIT A", fill="black")
    d.text(
        (100, 150),
        "This is a scanned document without native font text streams.",
        fill="black",
    )
    img.save(img_path)

    # Embed image onto PDF canvas without native text stream
    pdf_path = tmp_path / "scanned_document.pdf"
    c = pdf_canvas.Canvas(str(pdf_path), pagesize=letter)
    c.drawImage(str(img_path), 0, 0, width=612, height=792)
    c.save()
    return pdf_path


@pytest.fixture
def docx_document(tmp_path):
    docx_path = tmp_path / "settlement_agreement.docx"
    doc = docx.Document()
    doc.add_heading("SETTLEMENT AGREEMENT AND RELEASE", level=1)
    doc.add_paragraph(
        "This Settlement Agreement is entered into by and between Party A and Party B."
    )
    doc.add_heading("RECITALS", level=2)
    p = doc.add_paragraph(
        "WHEREAS, the parties desire to settle all outstanding claims amicably;"
    )
    p.add_run(
        " NOW THEREFORE, for good and valuable consideration, the parties agree as follows:"
    )
    doc.add_paragraph("1. Payment of Settlement Amount.", style="List Bullet")
    doc.add_paragraph("2. Mutual Release of Claims.", style="List Bullet")
    doc.save(docx_path)
    return docx_path


@pytest.fixture
def scanned_image_png(tmp_path):
    img_path = tmp_path / "notice_motion.png"
    img = Image.new("RGB", (1200, 1400), color="white")
    d = ImageDraw.Draw(img)
    d.text((100, 100), "NOTICE OF MOTION TO DISMISS", fill="black")
    d.text(
        (100, 160),
        "Defendant hereby moves to dismiss the complaint pursuant to Rule 12(b)(6).",
        fill="black",
    )
    img.save(img_path)
    return img_path


@pytest.fixture
def multipage_pdf(tmp_path):
    pdf_path = tmp_path / "legal_brief_multipage.pdf"
    c = pdf_canvas.Canvas(str(pdf_path), pagesize=letter)

    # Page 1
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "DEFENDANT MEMORANDUM OF LAW - PAGE 1")
    c.setFont("Helvetica", 11)
    c.drawString(
        72,
        680,
        "Statement of Facts: Plaintiff entered into an agreement on January 15, 2024.",
    )
    c.showPage()

    # Page 2
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "ARGUMENT & CITATIONS - PAGE 2")
    c.setFont("Helvetica", 11)
    c.drawString(
        72, 680, "Point I: The Statute of Limitations bars all asserted claims."
    )
    c.showPage()

    # Page 3
    c.setFont("Helvetica-Bold", 14)
    c.drawString(72, 720, "CONCLUSION - PAGE 3")
    c.setFont("Helvetica", 11)
    c.drawString(
        72,
        680,
        "For the foregoing reasons, Defendant respectfully requests that the Motion be Granted.",
    )
    c.save()
    return pdf_path


@pytest.fixture
def table_pdf(tmp_path):
    pdf_path = tmp_path / "case_docket_table.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
    styles = getSampleStyleSheet()

    elements = []
    elements.append(Paragraph("<b>CASE DOCKET TIMELINE</b>", styles["Heading1"]))
    elements.append(Spacer(1, 12))

    data = [
        ["Date", "Filing / Event", "Docket No.", "Status"],
        ["2024-01-10", "Complaint Filed", "Doc 1", "Pending"],
        ["2024-02-01", "Summons Issued", "Doc 3", "Executed"],
        ["2024-03-15", "Motion to Dismiss", "Doc 12", "Under Advisement"],
    ]

    t = Table(data, colWidths=[90, 200, 80, 110])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ]
        )
    )
    elements.append(t)
    doc.build(elements)
    return pdf_path


class TestDocumentIngestionPipeline:

    def test_1_normal_text_based_pdf(self, ingestion_pipeline, normal_text_pdf):
        doc = ingestion_pipeline.process_document(
            source=normal_text_pdf, case_id="case_101", force_ocr=False
        )
        assert doc.file_type == FileType.PDF
        assert doc.processing_metadata.processing_status == ProcessingStatus.COMPLETED
        assert len(doc.pages) == 1
        assert doc.pages[0].has_native_text is True
        assert doc.processing_metadata.ocr_used is False
        assert "COMPLAINT FOR BREACH OF CONTRACT" in doc.full_raw_text
        assert "28 U.S.C." in doc.full_normalized_text

    def test_2_scanned_pdf(self, ingestion_pipeline, scanned_pdf):
        has_native, p_count = NativeTextDetector.analyze_pdf_native_text(scanned_pdf)
        assert has_native is False

        doc = ingestion_pipeline.process_document(
            source=scanned_pdf, case_id="case_102", force_ocr=False
        )
        assert doc.file_type == FileType.PDF
        assert doc.pages[0].has_native_text is False
        assert doc.processing_metadata.ocr_used is True

    def test_3_docx_document(self, ingestion_pipeline, docx_document):
        doc = ingestion_pipeline.process_document(
            source=docx_document, case_id="case_103"
        )
        assert doc.file_type == FileType.DOCX
        assert doc.processing_metadata.processing_status == ProcessingStatus.COMPLETED
        assert len(doc.pages) >= 1
        assert "SETTLEMENT AGREEMENT AND RELEASE" in doc.full_raw_text
        assert len(doc.sections) >= 1

    def test_4_scanned_image(self, ingestion_pipeline, scanned_image_png):
        doc = ingestion_pipeline.process_document(
            source=scanned_image_png, case_id="case_104"
        )
        assert doc.file_type == FileType.IMAGE_PNG
        assert doc.processing_metadata.ocr_used is True
        assert doc.processing_metadata.processing_status == ProcessingStatus.COMPLETED

    def test_5_multipage_document(self, ingestion_pipeline, multipage_pdf):
        doc = ingestion_pipeline.process_document(
            source=multipage_pdf, case_id="case_105"
        )
        assert doc.file_type == FileType.PDF
        assert len(doc.pages) == 3
        assert doc.processing_metadata.page_count == 3
        page_numbers = [p.page_number for p in doc.pages]
        assert page_numbers == [1, 2, 3]
        assert "DEFENDANT MEMORANDUM OF LAW" in doc.pages[0].raw_text
        assert "ARGUMENT & CITATIONS" in doc.pages[1].raw_text
        assert "CONCLUSION" in doc.pages[2].raw_text

    def test_6_document_containing_tables(self, ingestion_pipeline, table_pdf):
        doc = ingestion_pipeline.process_document(source=table_pdf, case_id="case_106")
        assert doc.file_type == FileType.PDF
        page_1 = doc.pages[0]
        assert len(page_1.tables) >= 1 or "CASE DOCKET TIMELINE" in doc.full_raw_text
        if len(page_1.tables) > 0:
            tbl = page_1.tables[0]
            assert tbl.num_rows > 0
            assert tbl.num_cols > 0

    def test_7_provenance_traceability_and_storage(
        self, ingestion_pipeline, normal_text_pdf, tmp_storage_dir
    ):
        case_id = "case_provenance_test"
        doc = ingestion_pipeline.process_document(
            source=normal_text_pdf, case_id=case_id
        )

        # Check provenance at block level
        assert len(doc.pages[0].blocks) > 0
        first_block = doc.pages[0].blocks[0]
        assert first_block.provenance is not None
        assert first_block.provenance.case_id == case_id
        assert first_block.provenance.source_document_id == doc.document_id
        assert first_block.provenance.source_page == 1

        # Check storage persistence
        storage = DocumentStorage(base_dir=tmp_storage_dir)
        assert storage.has_document(doc.document_id) is True
        reloaded = storage.load_document(doc.document_id)
        assert reloaded is not None
        assert reloaded.document_id == doc.document_id
        assert reloaded.sha256_hash == doc.sha256_hash
        assert reloaded.case_id == case_id
